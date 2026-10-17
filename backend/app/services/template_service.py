# backend/app/services/template_service.py
"""Servicio de gestión de plantillas DOCX: versionamiento, análisis de variables."""

from __future__ import annotations

import re
import uuid
from pathlib import Path

from docx import Document as DocxDocument
from docxtpl import DocxTemplate
from jinja2.exceptions import TemplateSyntaxError
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.template import Template, TemplateField, TemplateVersion
from app.models.user import User
from app.schemas.template import DetectedVariable
from app.services.audit_service import AuditAction, AuditService

JINJA_VAR_RE = re.compile(r"\{\{-?\s*([\w.]+)\s*-?\}\}")
# Marcadores asistidos: guiones bajos, corchetes o secuencias de guiones.
MARKER_PATTERNS = [
    re.compile(r"\[\[([\w.]+)\]\]"),
    re.compile(r"_{4,}"),
    re.compile(r"-{4,}"),
]


# Palabras clave que sugieren el tipo de campo a partir de la clave/variable.
_TYPE_HINTS: list[tuple[re.Pattern[str], str]] = [
    (re.compile(r"(^|_)(dpi)($|_)", re.I), "dpi"),
    (re.compile(r"(^|_)(nit)($|_)", re.I), "nit"),
    (re.compile(r"(^|_)(correo|email)($|_)", re.I), "email"),
    (re.compile(r"(^|_)(telefono|celular|phone)($|_)", re.I), "phone"),
    (re.compile(r"(^|_)(fecha)($|_)", re.I), "date"),
    (
        re.compile(r"(^|_)(precio|monto|renta|capital|saldo|valor|total|subtotal)($|_)", re.I),
        "currency",
    ),
    (re.compile(r"(^|_)(porcentaje|pct)($|_)", re.I), "percentage"),
    (re.compile(r"(^|_)(cantidad|numero|plazo|meses|edad)($|_)", re.I), "integer"),
    (re.compile(r"(^|_)(nombre|nombres|razon_social|apellidos)($|_)", re.I), "name"),
    (re.compile(r"(^|_)(observaciones|descripcion|direccion|clausula)($|_)", re.I), "textarea"),
    (re.compile(r"(^|_)(es_|activo|casado|boolean)", re.I), "boolean"),
]


class TemplateService:
    def __init__(self, db: Session) -> None:
        self.db = db
        self._audit = AuditService(db)
        self._base_uploads = Path(settings.UPLOADS_DIR) / "templates"
        self._base_uploads.mkdir(parents=True, exist_ok=True)

    def extract_variables(self, docx_path: str) -> list[str]:
        """Extrae variables {{ var }} de un archivo DOCX."""
        variables: set[str] = set()
        try:
            doc = DocxDocument(docx_path)
            for para in doc.paragraphs:
                variables.update(JINJA_VAR_RE.findall(para.text))
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        for para in cell.paragraphs:
                            variables.update(JINJA_VAR_RE.findall(para.text))
        except Exception:
            pass
        return sorted(variables)

    def validate_docx_syntax(self, docx_path: str) -> str | None:
        """Comprueba que los marcadores Jinja2 de la plantilla sean válidos.

        Reutiliza el compilador de docxtpl (``get_undeclared_template_variables``)
        sobre el cuerpo y los encabezados/pies sin renderizar con datos, por lo
        que detecta delimitadores ``{{ }}`` / ``{% %}`` mal formados sin ejecutar
        nada. Devuelve el mensaje de error o ``None`` si la sintaxis es correcta.
        """
        try:
            DocxTemplate(docx_path).get_undeclared_template_variables()
        except TemplateSyntaxError as exc:
            return exc.message
        except Exception:
            # Un archivo ilegible se valida en otro punto del flujo; aquí solo
            # interesa la sintaxis de las variables.
            return None
        return None

    def save_uploaded_docx(self, content: bytes, original_name: str) -> str:
        """Guarda el DOCX con nombre interno seguro. Devuelve la ruta."""
        safe_name = f"{uuid.uuid4().hex}_{uuid.uuid4().hex[:4]}.docx"
        path = self._base_uploads / safe_name
        path.write_bytes(content)
        return str(path)

    def create_version(
        self,
        template: Template,
        file_path: str,
        original_name: str,
        notes: str | None = None,
        user: User | None = None,
    ) -> TemplateVersion:
        """Crea una nueva versión de plantilla sin sobrescribir la anterior."""
        last_version = (
            self.db.query(TemplateVersion)
            .filter(TemplateVersion.template_id == template.id)
            .order_by(TemplateVersion.version_number.desc())
            .first()
        )
        next_number = (last_version.version_number + 1) if last_version else 1

        variables = self.extract_variables(file_path)

        version = TemplateVersion(
            template_id=template.id,
            version_number=next_number,
            is_active=False,
            file_path=file_path,
            file_original_name=original_name,
            variables_detected=variables,
            notes=notes,
            created_by=user.id if user else None,
        )
        self.db.add(version)
        self.db.flush()
        self._audit.log(
            action=AuditAction.CREATE_VERSION,
            module="templates",
            user=user,
            record_id=version.id,
            detail=f"template={template.id} v{next_number}",
        )
        return version

    def activate_version(
        self, version: TemplateVersion, user: User | None = None
    ) -> TemplateVersion:
        """Desactiva todas las versiones y activa la indicada."""
        self.db.query(TemplateVersion).filter(
            TemplateVersion.template_id == version.template_id,
            TemplateVersion.is_active == True,  # noqa: E712
        ).update({"is_active": False})
        version.is_active = True
        self.db.flush()
        return version

    # ------------------------------------------------------------------
    # Análisis de contenido (mecanismo primario + asistido)
    # ------------------------------------------------------------------
    def analyze_docx(self, docx_path: str) -> tuple[list[DetectedVariable], list[DetectedVariable]]:
        """Devuelve (variables explícitas, marcadores asistidos)."""
        variables: list[DetectedVariable] = []
        markers: list[DetectedVariable] = []
        seen_vars: set[str] = set()
        seen_markers: set[str] = set()

        def scan(text: str, container: bool = False) -> None:
            for match in JINJA_VAR_RE.findall(text or ""):
                key = match.split(".")[-1]
                if match not in seen_vars:
                    seen_vars.add(match)
                    variables.append(
                        self._suggest(match, key, is_container=container or "." in match)
                    )
            for pattern in MARKER_PATTERNS:
                for match in pattern.findall(text or ""):
                    token = match if isinstance(match, str) and match else pattern.pattern
                    if token not in seen_markers:
                        seen_markers.add(token)
                        markers.append(
                            DetectedVariable(
                                variable=token,
                                suggested_key=f"marcador_{len(seen_markers)}",
                                suggested_label=f"Marcador {len(seen_markers)}",
                                suggested_field_type="text",
                            )
                        )

        try:
            doc = DocxDocument(docx_path)
        except Exception:
            return variables, markers

        for para in doc.paragraphs:
            scan(para.text)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        scan(para.text)

        return variables, markers

    def _suggest(self, variable: str, key: str, *, is_container: bool = False) -> DetectedVariable:
        """Sugiere metadatos para una variable detectada (no definitivos)."""
        if is_container:
            field_type = "list"
        else:
            field_type = "text"
            for pattern, candidate in _TYPE_HINTS:
                if pattern.search(key):
                    field_type = candidate
                    break
        label = key.replace("_", " ").strip().title()
        return DetectedVariable(
            variable=variable,
            suggested_key=key,
            suggested_label=label,
            suggested_field_type=field_type,
            is_container=is_container,
        )

    def create_fields_from_variables(
        self,
        version: TemplateVersion,
        variables: list[str],
        *,
        replace_existing: bool = False,
    ) -> list[TemplateField]:
        """Crea definiciones de campo a partir de variables confirmadas por el usuario."""
        if replace_existing:
            self.db.query(TemplateField).filter(
                TemplateField.template_version_id == version.id
            ).delete()

        existing = {f.key for f in version.fields}
        created: list[TemplateField] = []
        order = len(existing)
        for variable in variables:
            key = variable.split(".")[-1]
            if key in existing:
                continue
            suggestion = self._suggest(variable, key, is_container="." in variable)
            field = TemplateField(
                template_version_id=version.id,
                key=key,
                label=suggestion.suggested_label,
                field_type=suggestion.suggested_field_type,
                required=False,
                nullable=True,
                docx_variable=variable,
                display_order=order,
                active=True,
            )
            self.db.add(field)
            created.append(field)
            existing.add(key)
            order += 1
        self.db.flush()
        return created
