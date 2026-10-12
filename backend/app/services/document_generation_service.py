# backend/app/services/document_generation_service.py
"""Servicio de generación DOCX usando docxtpl."""

from __future__ import annotations

import re
import uuid
from pathlib import Path

from docx import Document as DocxDocument
from docx.document import Document as DocxDocumentType
from docxtpl import DocxTemplate
from sqlalchemy.orm import Session

from app.core.config import settings
from app.models.document import Document, DocumentVersion
from app.models.user import User
from app.services.audit_service import AuditAction, AuditService


class DocumentGenerationError(Exception):
    """Error de generación de documento."""


class DocumentGenerationService:
    """Genera archivos DOCX desde plantillas docxtpl."""

    TEMPLATE_VAR_RE = re.compile(r"\{\{[^{]*?\}\}")

    def __init__(self, db: Session) -> None:
        self.db = db
        self._audit = AuditService(db)
        self._base_generated = Path(settings.GENERATED_DIR) / "documents"
        self._base_generated.mkdir(parents=True, exist_ok=True)

    def generate(self, document: Document, user: User | None = None) -> DocumentVersion:
        """Genera el DOCX y crea una nueva DocumentVersion."""
        if document.template_version is None:
            raise DocumentGenerationError("El documento no tiene plantilla asignada.")

        template_path = document.template_version.file_path
        if not template_path or not Path(template_path).exists():
            raise DocumentGenerationError(f"Archivo de plantilla no encontrado: {template_path}")

        context = document.field_data or {}

        # Calcular versión
        version_number = (
            self.db.query(DocumentVersion)
            .filter(DocumentVersion.document_id == document.id)
            .count()
        ) + 1

        # Nombre de archivo seguro
        safe_name = f"doc_{document.id}_{version_number}_{uuid.uuid4().hex[:8]}.docx"
        output_path = self._base_generated / safe_name

        # Renderizar con docxtpl
        try:
            tpl = DocxTemplate(template_path)
            tpl.render(context)
            tpl.save(str(output_path))
        except Exception as e:
            raise DocumentGenerationError(f"Error al renderizar plantilla: {e}") from e

        # Verificar que el archivo existe
        if not output_path.exists():
            raise DocumentGenerationError(
                "El archivo generado no fue encontrado después de la generación."
            )

        # Verificar que se puede abrir con python-docx
        try:
            doc_check = DocxDocument(str(output_path))
        except Exception as e:
            raise DocumentGenerationError(f"El DOCX generado no se puede abrir: {e}") from e

        # Verificar que no quedan variables sin sustituir
        unresolved = self._find_unresolved(doc_check)
        if unresolved:
            raise DocumentGenerationError(
                f"El documento generado contiene variables sin sustituir: {unresolved}"
            )

        # Crear DocumentVersion
        dv = DocumentVersion(
            document_id=document.id,
            version_number=version_number,
            file_path=str(output_path),
            source_template_version_id=document.template_version_id,
            data_snapshot=context,
            validation_status="GENERATED",
            created_by=user.id if user else None,
        )
        self.db.add(dv)
        self.db.flush()

        self._audit.log(
            action=AuditAction.GENERATE_DOCUMENT,
            module="documents",
            user=user,
            record_id=document.id,
            detail=f"version={version_number} file={safe_name}",
        )
        return dv

    def _find_unresolved(self, doc: DocxDocumentType) -> list[str]:
        """Busca variables {{ }} sin sustituir en el DOCX."""
        unresolved = []
        for para in doc.paragraphs:
            matches = self.TEMPLATE_VAR_RE.findall(para.text)
            unresolved.extend(matches)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for para in cell.paragraphs:
                        matches = self.TEMPLATE_VAR_RE.findall(para.text)
                        unresolved.extend(matches)
        return list(set(unresolved))
