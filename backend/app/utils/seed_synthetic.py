# backend/app/utils/seed_synthetic.py
"""Generador de datos sintéticos y de referencia.

Uso (desde la carpeta `backend/`):

    python -m app.utils.seed_synthetic
    python -m app.utils.seed_synthetic --reset

Crea:
- roles, permisos y usuarios demo;
- tipos de escritura;
- 5 plantillas DOCX sintéticas con variables;
- 40 clientes sintéticos;
- 100 casos (20 por tipo) con inconsistencias controladas;
- documentos asociados y mediciones experimentales;
- una cotización, un presupuesto, un cobro y un pago de ejemplo.

Todos los datos son claramente sintéticos. No se usan datos personales reales.
"""

from __future__ import annotations

import argparse
import random
from datetime import UTC, date, datetime, timedelta
from decimal import Decimal
from pathlib import Path
from typing import Any

from docx import Document as DocxDocument
from faker import Faker
from sqlalchemy.orm import Session

from app.core.config import settings
from app.core.security import hash_password
from app.db.base import Base
from app.db.session import SessionLocal, engine
from app.models import (  # noqa: F401
    AuditLog,
    Budget,
    BudgetItem,
    Case,
    CaseParty,
    Charge,
    Client,
    DeedType,
    Document,
    LegalEntity,
    Payment,
    Permission,
    Person,
    Quote,
    QuoteItem,
    Role,
    Template,
    TemplateField,
    TemplateVersion,
    TestCase,
    TestExecution,
    User,
    ValidationResult,
    ValidationRule,
)
from app.models.measurement import TimeMeasurement
from app.rules.catalog import rule_applies
from app.utils.files import ensure_storage_dirs

fake = Faker("es_ES")
Faker.seed(20260101)
random.seed(20260101)

# ---------------------------------------------------------------------------
# Catálogos
# ---------------------------------------------------------------------------

PERMISSIONS: list[tuple[str, str, str]] = [
    ("users.read", "users", "Ver usuarios"),
    ("users.create", "users", "Crear usuarios"),
    ("users.update", "users", "Editar usuarios"),
    ("users.delete", "users", "Eliminar usuarios"),
    ("clients.read", "clients", "Ver clientes"),
    ("clients.create", "clients", "Crear clientes"),
    ("clients.update", "clients", "Editar clientes"),
    ("clients.delete", "clients", "Eliminar clientes"),
    ("cases.read", "cases", "Ver expedientes"),
    ("cases.create", "cases", "Crear expedientes"),
    ("cases.update", "cases", "Editar expedientes"),
    ("cases.delete", "cases", "Eliminar expedientes"),
    ("templates.read", "templates", "Ver plantillas"),
    ("templates.create", "templates", "Crear plantillas"),
    ("templates.update", "templates", "Editar plantillas"),
    ("templates.delete", "templates", "Eliminar plantillas"),
    ("templates.activate", "templates", "Activar versiones"),
    ("documents.read", "documents", "Ver documentos"),
    ("documents.create", "documents", "Crear documentos"),
    ("documents.update", "documents", "Editar documentos"),
    ("documents.generate", "documents", "Generar DOCX"),
    ("validations.read", "validations", "Ver validaciones"),
    ("validations.execute", "validations", "Ejecutar validaciones"),
    ("quotes.read", "quotes", "Ver cotizaciones"),
    ("quotes.create", "quotes", "Crear cotizaciones"),
    ("quotes.update", "quotes", "Editar cotizaciones"),
    ("quotes.delete", "quotes", "Eliminar cotizaciones"),
    ("budgets.read", "budgets", "Ver presupuestos"),
    ("budgets.create", "budgets", "Crear presupuestos"),
    ("budgets.update", "budgets", "Editar presupuestos"),
    ("charges.read", "charges", "Ver cobros"),
    ("charges.create", "charges", "Crear cobros"),
    ("charges.update", "charges", "Editar cobros"),
    ("payments.read", "payments", "Ver pagos"),
    ("payments.create", "payments", "Registrar pagos"),
    ("reports.read", "reports", "Ver reportes y dashboard"),
    ("audit.read", "audit", "Ver auditoría"),
    ("rules.read", "rules", "Ver reglas de validación"),
    ("rules.update", "rules", "Editar reglas de validación"),
    ("deed_types.create", "deed_types", "Crear tipos de escritura"),
    ("deed_types.update", "deed_types", "Editar tipos de escritura"),
]

ROLE_DEFS: dict[str, list[str]] = {
    "ADMINISTRADOR": [p[0] for p in PERMISSIONS],
    "ABOGADO_NOTARIO": [
        "clients.read",
        "clients.create",
        "clients.update",
        "clients.delete",
        "cases.read",
        "cases.create",
        "cases.update",
        "cases.delete",
        "templates.read",
        "templates.create",
        "templates.update",
        "templates.activate",
        "documents.read",
        "documents.create",
        "documents.update",
        "documents.generate",
        "validations.read",
        "validations.execute",
        "quotes.read",
        "quotes.create",
        "quotes.update",
        "budgets.read",
        "budgets.create",
        "budgets.update",
        "charges.read",
        "charges.create",
        "charges.update",
        "payments.read",
        "payments.create",
        "reports.read",
    ],
    "AUXILIAR": [
        "clients.read",
        "clients.create",
        "clients.update",
        "cases.read",
        "cases.create",
        "cases.update",
        "templates.read",
        "documents.read",
        "documents.create",
        "documents.update",
        "validations.read",
        "validations.execute",
        "quotes.read",
        "charges.read",
        "reports.read",
    ],
    "ADMINISTRACION": [
        "clients.read",
        "cases.read",
        "templates.read",
        "documents.read",
        "quotes.read",
        "quotes.create",
        "quotes.update",
        "budgets.read",
        "budgets.create",
        "budgets.update",
        "charges.read",
        "charges.create",
        "charges.update",
        "payments.read",
        "payments.create",
        "reports.read",
    ],
}

DEED_TYPES: list[tuple[str, str, str]] = [
    (
        "COMPRAVENTA_INMUEBLE",
        "Compraventa de inmueble",
        "Traslado de dominio de un inmueble entre vendedor y comprador.",
    ),
    (
        "DONACION_ENTRE_VIVOS",
        "Donación entre vivos",
        "Donación de bienes del donante al donatario.",
    ),
    (
        "ARRENDAMIENTO",
        "Arrendamiento",
        "Contrato de arrendamiento de inmueble por plazo determinado.",
    ),
    (
        "PROTOCOLACION_MATRIMONIO",
        "Protocolación de matrimonio",
        "Protocolización notarial del matrimonio civil.",
    ),
    (
        "CONSTITUCION_SOCIEDAD",
        "Constitución de sociedad",
        "Constitución de sociedad anónima con accionistas y aportes.",
    ),
]

# Definición de plantillas: variables y cuerpo del DOCX.
TEMPLATE_DEFS: dict[str, dict] = {
    "COMPRAVENTA_INMUEBLE": {
        "name": "Compraventa de inmueble (sintética)",
        "fields": [
            ("numero_escritura", "Número de escritura", "integer", True),
            ("fecha_escritura", "Fecha de escritura", "date", True),
            ("comprador_nombre", "Comprador", "name", True),
            ("comprador_dpi", "DPI del comprador", "dpi", True),
            ("comprador_nit", "NIT del comprador", "nit", False),
            ("vendedor_nombre", "Vendedor", "name", True),
            ("vendedor_dpi", "DPI del vendedor", "dpi", True),
            ("vendedor_nit", "NIT del vendedor", "nit", False),
            ("precio", "Precio", "currency", True),
            ("finca", "Finca", "text", True),
            ("folio", "Folio", "text", True),
            ("libro", "Libro", "text", True),
            ("departamento", "Departamento", "select", True),
            ("municipio", "Municipio", "text", True),
        ],
        "body": [
            "ESCRITURA PÚBLICA NÚMERO {{ numero_escritura }}.",
            "En la ciudad de Guatemala, el {{ fecha_escritura }}, ante mí, NOTARIO, comparecen:",
            "Por una parte, el señor(es) {{ vendedor_nombre }}, con DPI "
            "{{ vendedor_dpi }}{% if vendedor_nit %} y NIT {{ vendedor_nit }}{% endif %}, "
            "a quien en lo sucesivo se denominará VENDEDOR.",
            "Por otra parte, el señor(es) {{ comprador_nombre }}, con DPI "
            "{{ comprador_dpi }}{% if comprador_nit %} y NIT {{ comprador_nit }}{% endif %}, "
            "a quien en lo sucesivo se denominará COMPRADOR.",
            "PRIMERA. El VENDEDOR transfiere el inmueble identificado con finca "
            "número {{ finca }}, folio {{ folio }}, libro {{ libro }}, ubicado en "
            "el departamento de {{ departamento }}, municipio de {{ municipio }}.",
            "SEGUNDA. El precio de la compraventa es de {{ precio }}, que el "
            "COMPRADOR paga en este acto.",
            "TERCERA. El VENDEDOR declara que el inmueble se encuentra libre de "
            "gravámenes y anotaciones.",
        ],
    },
    "DONACION_ENTRE_VIVOS": {
        "name": "Donación entre vivos (sintética)",
        "fields": [
            ("numero_escritura", "Número de escritura", "integer", True),
            ("fecha_escritura", "Fecha de escritura", "date", True),
            ("donante_nombre", "Donante", "name", True),
            ("donante_dpi", "DPI del donante", "dpi", True),
            ("donatario_nombre", "Donatario", "name", True),
            ("donatario_dpi", "DPI del donatario", "dpi", True),
            ("bienes", "Bienes donados", "textarea", True),
            ("finca", "Finca", "text", True),
            ("folio", "Folio", "text", True),
            ("libro", "Libro", "text", True),
        ],
        "body": [
            "ESCRITURA PÚBLICA NÚMERO {{ numero_escritura }}.",
            "En la ciudad de Guatemala, el {{ fecha_escritura }}, ante mí, NOTARIO, comparecen:",
            "El señor(es) {{ donante_nombre }}, con DPI {{ donante_dpi }}, a quien "
            "se denominará DONANTE.",
            "El señor(es) {{ donatario_nombre }}, con DPI {{ donatario_dpi }}, a "
            "quien se denominará DONATARIO.",
            "PRIMERA. El DONANTE dona de forma gratuita e irrevocable los siguientes "
            "bienes: {{ bienes }}.",
            "SEGUNDA. El bien se identifica con finca {{ finca }}, folio {{ folio }}, "
            "libro {{ libro }}.",
        ],
    },
    "ARRENDAMIENTO": {
        "name": "Arrendamiento (sintética)",
        "fields": [
            ("numero_escritura", "Número de escritura", "integer", True),
            ("fecha_escritura", "Fecha de escritura", "date", True),
            ("arrendador_nombre", "Arrendador", "name", True),
            ("arrendador_dpi", "DPI del arrendador", "dpi", True),
            ("arrendatario_nombre", "Arrendatario", "name", True),
            ("arrendatario_dpi", "DPI del arrendatario", "dpi", True),
            ("inmueble", "Inmueble", "textarea", True),
            ("renta", "Renta mensual", "currency", True),
            ("plazo_meses", "Plazo (meses)", "integer", True),
            ("fecha_inicio", "Fecha de inicio", "date", True),
            ("fecha_fin", "Fecha de finalización", "date", True),
        ],
        "body": [
            "ESCRITURA PÚBLICA NÚMERO {{ numero_escritura }}.",
            "En la ciudad de Guatemala, el {{ fecha_escritura }}, comparecen "
            "{{ arrendador_nombre }} con DPI {{ arrendador_dpi }}, ARRENDADOR, y "
            "{{ arrendatario_nombre }} con DPI {{ arrendatario_dpi }}, ARRENDATARIO.",
            "PRIMERA. El ARRENDADOR entrega en arrendamiento el inmueble: {{ inmueble }}.",
            "SEGUNDA. La renta mensual es de {{ renta }}, pagadera dentro de los "
            "primeros cinco días de cada mes.",
            "TERCERA. El plazo del arrendamiento es de {{ plazo_meses }} meses, "
            "iniciando el {{ fecha_inicio }} y finalizando el {{ fecha_fin }}.",
        ],
    },
    "PROTOCOLACION_MATRIMONIO": {
        "name": "Protocolación de matrimonio (sintética)",
        "fields": [
            ("numero_escritura", "Número de escritura", "integer", True),
            ("fecha_escritura", "Fecha de escritura", "date", True),
            ("compareciente1_nombre", "Primer compareciente", "name", True),
            ("compareciente1_dpi", "DPI del primer compareciente", "dpi", True),
            ("compareciente2_nombre", "Segundo compareciente", "name", True),
            ("compareciente2_dpi", "DPI del segundo compareciente", "dpi", True),
            ("fecha_matrimonio", "Fecha de matrimonio", "date", True),
            ("referencia_acta", "Referencia del acta", "text", True),
        ],
        "body": [
            "ESCRITURA PÚBLICA NÚMERO {{ numero_escritura }}.",
            "En la ciudad de Guatemala, el {{ fecha_escritura }}, comparecen "
            "{{ compareciente1_nombre }} con DPI {{ compareciente1_dpi }} y "
            "{{ compareciente2_nombre }} con DPI {{ compareciente2_dpi }}.",
            "PRIMERA. Los comparecientes solicitan la protocolización de su "
            "matrimonio civil celebrado el {{ fecha_matrimonio }}.",
            "SEGUNDA. El acta de matrimonio se identifica con la referencia {{ referencia_acta }}.",
        ],
    },
    "CONSTITUCION_SOCIEDAD": {
        "name": "Constitución de sociedad (sintética)",
        "fields": [
            ("numero_escritura", "Número de escritura", "integer", True),
            ("fecha_escritura", "Fecha de escritura", "date", True),
            ("razon_social", "Razón social", "text", True),
            ("nombre_comercial", "Nombre comercial", "text", False),
            ("capital_inicial", "Capital inicial", "currency", True),
            ("representante_legal", "Representante legal", "name", True),
            ("accionistas", "Accionistas", "list", True),
        ],
        "body": [
            "ESCRITURA PÚBLICA NÚMERO {{ numero_escritura }}.",
            "En la ciudad de Guatemala, el {{ fecha_escritura }}, comparecen los "
            "señores accionistas para constituir la sociedad {{ razon_social }}.",
            "PRIMERA. La sociedad girará bajo la razón social {{ razon_social }}"
            "{% if nombre_comercial %} y nombre comercial {{ nombre_comercial }}{% endif %}.",
            "SEGUNDA. El capital social inicial es de {{ capital_inicial }}.",
            "TERCERA. Son accionistas fundadores: "
            "{% for accionista in accionistas %}{{ accionista.nombre }} "
            "({{ accionista.aporte }}); {% endfor %}.",
            "CUARTA. Se designa como representante legal a {{ representante_legal }}.",
        ],
    },
}

INCONSISTENCY_TYPES = [
    "VALIDEZ_PLENA",
    "DPI_INCONSISTENTE",
    "NIT_INCONSISTENTE",
    "NOMBRE_INCONSISTENTE",
    "FECHA_INCONSISTENTE",
    "MONTO_INCONSISTENTE",
    "REGISTRAL_INCONSISTENTE",
    "INCISO_FALTANTE",
    "INCISO_DUPLICADO",
    "CAMPO_OBLIGATORIO_VACIO",
]

DEPARTAMENTOS = [
    "Guatemala",
    "Sacatepéquez",
    "Chimaltenango",
    "Escuintla",
    "Quetzaltenango",
    "Huehuetenango",
    "Petén",
    "Izabal",
    "Zacapa",
    "Suchitepéquez",
]
MUNICIPIOS = [
    "Guatemala",
    "Mixco",
    "Villa Nueva",
    "Antigua Guatemala",
    "Quetzaltenango",
    "Cobán",
    "Flores",
    "Puerto Barrios",
    "Zacapa",
    "Mazatenango",
]


# ---------------------------------------------------------------------------
# DOCX
# ---------------------------------------------------------------------------


def build_template_docx(path: Path, title: str, body: list[str]) -> None:
    """Construye un DOCX sintético con marcadores docxtpl."""
    doc = DocxDocument()
    doc.add_heading(title, level=1)
    doc.add_paragraph("Documento académico sintético — no tiene validez legal.")
    for line in body:
        doc.add_paragraph(line)
    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(path))


# ---------------------------------------------------------------------------
# Utilidades de datos sintéticos
# ---------------------------------------------------------------------------


def synthetic_dpi() -> str:
    return "".join(str(random.randint(0, 9)) for _ in range(13))


def synthetic_nit() -> str:
    base = "".join(str(random.randint(0, 9)) for _ in range(8))
    check = random.choice(list("0123456789") + ["K"])
    return f"{base}-{check}"


def synthetic_name() -> str:
    return f"{fake.first_name()} {fake.last_name()} {fake.last_name()}"


def synthetic_phone() -> str:
    return f"{random.randint(3, 7)}{random.randint(10_000_000, 99_999_999)}"


def synthetic_email(name: str) -> str:
    slug = name.lower().replace(" ", ".").replace("á", "a").replace("é", "e")
    slug = slug.replace("í", "i").replace("ó", "o").replace("ú", "u").replace("ñ", "n")
    return f"{slug}@example.local"


# ---------------------------------------------------------------------------
# Datos de referencia
# ---------------------------------------------------------------------------


def seed_reference_data(db: Session) -> dict[str, Any]:
    roles = {}
    for role_name, _perm_codes in ROLE_DEFS.items():
        role = db.query(Role).filter(Role.name == role_name).first()
        if not role:
            role = Role(name=role_name, description=f"Rol {role_name}")
            db.add(role)
            db.flush()
        roles[role_name] = role

    permissions = {}
    for code, module, desc in PERMISSIONS:
        perm = db.query(Permission).filter(Permission.code == code).first()
        if not perm:
            perm = Permission(code=code, module=module, description=desc)
            db.add(perm)
            db.flush()
        permissions[code] = perm

    for role_name, perm_codes in ROLE_DEFS.items():
        role = roles[role_name]
        role.permissions = [permissions[c] for c in perm_codes]
    db.flush()

    deed_types = {}
    for code, name, description in DEED_TYPES:
        dt = db.query(DeedType).filter(DeedType.code == code).first()
        if not dt:
            dt = DeedType(code=code, name=name, description=description)
            db.add(dt)
            db.flush()
        deed_types[code] = dt
    db.flush()
    return {"roles": roles, "permissions": permissions, "deed_types": deed_types}


def seed_validation_rules(db: Session) -> None:
    """Crea las reglas de validación administrables a partir del catálogo.

    Es idempotente: solo inserta las reglas que falten, respetando cualquier
    personalización previa del administrador.
    """

    from app.services.rule_service import RuleService

    RuleService(db).ensure_rules()
    db.flush()


def seed_users(db: Session, roles: dict) -> dict[str, User]:
    demo = [
        ("admin", "admin@example.local", "Administrador Demo", "Admin123!", "ADMINISTRADOR"),
        ("notario.demo", "notario@example.local", "Notario Demo", "Notario123!", "ABOGADO_NOTARIO"),
        ("auxiliar.demo", "auxiliar@example.local", "Auxiliar Demo", "Auxiliar123!", "AUXILIAR"),
        (
            "adminfin.demo",
            "adminfin@example.local",
            "Admin Financiero Demo",
            "Adminfin123!",
            "ADMINISTRACION",
        ),
    ]
    users = {}
    for username, email, full_name, password, role_name in demo:
        user = db.query(User).filter(User.username == username).first()
        if not user:
            user = User(
                username=username,
                email=email,
                full_name=full_name,
                password_hash=hash_password(password),
                role_id=roles[role_name].id,
                status="ACTIVE",
            )
            db.add(user)
            db.flush()
        users[username] = user
    return users


def seed_templates(db: Session, deed_types: dict, users: dict) -> dict[str, Template]:
    templates: dict[str, Template] = {}
    base = Path(settings.UPLOADS_DIR) / "templates"
    base.mkdir(parents=True, exist_ok=True)
    creator = users["notario.demo"]

    for code, definition in TEMPLATE_DEFS.items():
        template = db.query(Template).filter(Template.name == definition["name"]).first()
        if template:
            templates[code] = template
            continue

        template = Template(
            name=definition["name"],
            description=f"Plantilla académica sintética para {deed_types[code].name}.",
            deed_type_id=deed_types[code].id,
            created_by=creator.id,
        )
        db.add(template)
        db.flush()

        file_name = f"plantilla_{code.lower()}_v1.docx"
        file_path = base / file_name
        build_template_docx(file_path, definition["name"], definition["body"])

        variables = []
        for key, _label, _tipo, _req in definition["fields"]:
            variables.append(key)

        version = TemplateVersion(
            template_id=template.id,
            version_number=1,
            is_active=True,
            file_path=str(file_path),
            file_original_name=file_name,
            variables_detected=variables,
            notes="Versión inicial generada por datos sintéticos.",
            created_by=creator.id,
        )
        db.add(version)
        db.flush()

        for order, (key, label, ftype, required) in enumerate(definition["fields"]):
            options = None
            if ftype == "select" and key == "departamento":
                options = [
                    {"label": d, "value": d, "active": True, "order": i}
                    for i, d in enumerate(DEPARTAMENTOS)
                ]
            db.add(
                TemplateField(
                    template_version_id=version.id,
                    key=key,
                    label=label,
                    field_type=ftype,
                    required=required,
                    nullable=not required,
                    docx_variable=key,
                    display_order=order,
                    options_json=options,
                    help_text=None,
                    active=True,
                )
            )
        db.flush()
        templates[code] = template
    db.flush()
    return templates


def seed_clients(db: Session, count: int = 40) -> list[Client]:
    clients: list[Client] = []
    existing = db.query(Client).count()
    if existing:
        return db.query(Client).all()

    for i in range(count):
        if i % 5 == 0:
            entity = LegalEntity(
                client_id="",
                razon_social=f"Sociedad Sintética {i:02d}, S.A.",
                nombre_comercial=f"SINTÉTICA {i:02d}",
                nit=synthetic_nit(),
                tipo_sociedad="S.A.",
                numero_inscripcion=f"REG-{1000 + i}",
                departamento=random.choice(DEPARTAMENTOS),
                municipio=random.choice(MUNICIPIOS),
                telefono=synthetic_phone(),
                correo=f"sociedad{i:02d}@example.local",
                representante_legal=synthetic_name(),
                cargo_representante="Gerente General",
            )
            client = Client(tipo_cliente="JURIDICA")
            db.add(client)
            db.flush()
            entity.client_id = client.id
            db.add(entity)
        else:
            nombre = synthetic_name()
            person = Person(
                client_id="",
                nombres=nombre.split()[0],
                apellidos=" ".join(nombre.split()[1:]),
                nombre_completo=nombre,
                dpi=synthetic_dpi(),
                nit=synthetic_nit(),
                nacionalidad="Guatemalteca",
                estado_civil=random.choice(["Soltero", "Casado", "Divorciado"]),
                profesion=random.choice(
                    ["Comerciante", "Ingeniero", "Abogado", "Estudiante", "Médico"]
                ),
                direccion=f"{fake.street_address()}",
                departamento=random.choice(DEPARTAMENTOS),
                municipio=random.choice(MUNICIPIOS),
                telefono=synthetic_phone(),
                correo=synthetic_email(nombre),
            )
            client = Client(tipo_cliente="INDIVIDUAL")
            db.add(client)
            db.flush()
            person.client_id = client.id
            db.add(person)
        clients.append(client)
    db.flush()
    return clients


def _person_for(db: Session, client: Client) -> Person:
    person = db.query(Person).filter(Person.client_id == client.id).first()
    assert person is not None
    return person


def _legal_for(db: Session, client: Client) -> LegalEntity:
    entity = db.query(LegalEntity).filter(LegalEntity.client_id == client.id).first()
    assert entity is not None
    return entity


# ---------------------------------------------------------------------------
# Documentos por tipo
# ---------------------------------------------------------------------------


def build_field_data(code: str, db: Session, clients: list[Client], numero: int) -> dict:
    """Construye field_data para un caso, incluyendo datos de expediente para consistencia."""
    people = [c for c in clients if c.tipo_cliente == "INDIVIDUAL"]
    a = _person_for(db, random.choice(people))
    b = _person_for(db, random.choice(people))
    c = _person_for(db, random.choice(people))
    fecha = date(2026, random.randint(1, 12), random.randint(1, 28))

    common = {
        "numero_escritura": numero,
        "fecha_escritura": fecha.isoformat(),
    }

    if code == "COMPRAVENTA_INMUEBLE":
        common.update(
            {
                "comprador_nombre": a.nombre_completo,
                "comprador_dpi": a.dpi,
                "comprador_nit": a.nit,
                "vendedor_nombre": b.nombre_completo,
                "vendedor_dpi": b.dpi,
                "vendedor_nit": b.nit,
                "precio": f"{random.randint(2, 20) * 100000}.00",
                "finca": f"{random.randint(1000, 99999)}",
                "folio": f"{random.randint(1, 500)}",
                "libro": f"{random.randint(1, 300)}",
                "departamento": random.choice(DEPARTAMENTOS),
                "municipio": random.choice(MUNICIPIOS),
                # Datos de expediente (para validar consistencia)
                "expediente_comprador_dpi": a.dpi,
                "expediente_comprador_nit": a.nit,
                "expediente_comprador_nombre": a.nombre_completo,
                "subtotal": f"{random.randint(2, 20) * 100000}.00",
            }
        )
        common["total"] = common["precio"]
    elif code == "DONACION_ENTRE_VIVOS":
        common.update(
            {
                "donante_nombre": a.nombre_completo,
                "donante_dpi": a.dpi,
                "donatario_nombre": b.nombre_completo,
                "donatario_dpi": b.dpi,
                "bienes": "Un lote de terreno y la casa en él construida, descritos de forma sintética.",
                "finca": f"{random.randint(1000, 99999)}",
                "folio": f"{random.randint(1, 500)}",
                "libro": f"{random.randint(1, 300)}",
            }
        )
    elif code == "ARRENDAMIENTO":
        inicio = fecha
        fin = inicio + timedelta(days=365)
        common.update(
            {
                "arrendador_nombre": a.nombre_completo,
                "arrendador_dpi": a.dpi,
                "arrendatario_nombre": b.nombre_completo,
                "arrendatario_dpi": b.dpi,
                "inmueble": "Inmueble sintético ubicado en zona 1, ciudad de Guatemala.",
                "renta": f"{random.randint(2500, 12000)}.00",
                "plazo_meses": 12,
                "fecha_inicio": inicio.isoformat(),
                "fecha_fin": fin.isoformat(),
            }
        )
    elif code == "PROTOCOLACION_MATRIMONIO":
        common.update(
            {
                "compareciente1_nombre": a.nombre_completo,
                "compareciente1_dpi": a.dpi,
                "compareciente2_nombre": b.nombre_completo,
                "compareciente2_dpi": b.dpi,
                "fecha_matrimonio": (fecha - timedelta(days=30)).isoformat(),
                "referencia_acta": f"ACTA-{random.randint(10000, 99999)}",
            }
        )
    elif code == "CONSTITUCION_SOCIEDAD":
        capital = random.randint(50, 500) * 1000
        aportes: list[dict[str, Any]] = []
        total = 0
        names = [a, b, c]
        for p in names:
            monto = capital // 3
            aportes.append({"nombre": p.nombre_completo, "aporte": f"{monto}.00", "monto": monto})
            total += monto
        # Ajustar residuo para que la suma sea exacta (evita RULE-018 falso positivo)
        diff = capital - total
        if diff:
            aportes[0]["monto"] += diff
            aportes[0]["aporte"] = f"{aportes[0]['monto']}.00"
            total += diff
        common.update(
            {
                "razon_social": f"Empresa Sintética {numero}, S.A.",
                "nombre_comercial": f"SINTÉTICA {numero}",
                "capital_inicial": f"{capital}.00",
                "representante_legal": a.nombre_completo,
                "accionistas": aportes,
                "aportes": aportes,
                "incisos": [
                    {"numero": "a", "texto": "Constituir la sociedad anónima."},
                    {"numero": "b", "texto": "Aprobar el capital social y su distribución."},
                    {"numero": "c", "texto": "Nombrar al representante legal."},
                ],
            }
        )
    return common


def _expected(kinds: list[str], code: str) -> int:
    """Cuenta las reglas aplicables al tipo de escritura para las inconsistencias."""

    return sum(1 for rule_id in kinds if rule_applies(rule_id, code))


def apply_inconsistency(code: str, data: dict, kind: str) -> int:
    """Introduce una inconsistencia controlada. Devuelve la cantidad esperada de errores.

    Solo se mutan los datos (y se cuenta el error) cuando la regla que detecta
    la inconsistencia aplica al tipo de escritura. Así, por ejemplo, una
    inconsistencia de incisos no afecta a una compraventa.
    """

    if kind == "VALIDEZ_PLENA":
        return 0
    if kind == "DPI_INCONSISTENTE":
        if not rule_applies("RULE-004", code):
            return 0
        data["comprador_dpi"] = synthetic_dpi()
        data["expediente_comprador_dpi"] = synthetic_dpi()
        return _expected(["RULE-004"], code)
    if kind == "NIT_INCONSISTENTE":
        if not rule_applies("RULE-005", code):
            return 0
        data["comprador_nit"] = synthetic_nit()
        data["expediente_comprador_nit"] = synthetic_nit()
        return _expected(["RULE-005"], code)
    if kind == "NOMBRE_INCONSISTENTE":
        if not rule_applies("RULE-006", code):
            return 0
        data["comprador_nombre"] = "Nombre Sintético Modificado"
        data["expediente_comprador_nombre"] = synthetic_name()
        return _expected(["RULE-006"], code)
    if kind == "FECHA_INCONSISTENTE":
        if code == "ARRENDAMIENTO":
            data["fecha_fin"] = "2024-01-01"
            data["fecha_inicio"] = "2026-01-01"
            return _expected(["RULE-007"], code)
        # Solo la inconsistencia de fechas del arrendamiento tiene regla propia.
        return 0
    if kind == "MONTO_INCONSISTENTE":
        if not rule_applies("RULE-008", code):
            return 0
        data["subtotal"] = "1000.00"
        data["total"] = "999999.00"
        return _expected(["RULE-008"], code)
    if kind == "REGISTRAL_INCONSISTENTE":
        if not rule_applies("RULE-009", code):
            return 0
        data["finca"] = ""
        data["folio"] = ""
        return _expected(["RULE-009", "RULE-010"], code)
    if kind == "INCISO_FALTANTE":
        if not rule_applies("RULE-014", code):
            return 0
        data["incisos"] = []
        return _expected(["RULE-014"], code)
    if kind == "INCISO_DUPLICADO":
        if not rule_applies("RULE-015", code):
            return 0
        data["incisos"] = [
            {"numero": "a", "texto": "Inciso uno"},
            {"numero": "a", "texto": "Inciso duplicado"},
        ]
        return _expected(["RULE-015"], code)
    if kind == "CAMPO_OBLIGATORIO_VACIO":
        if code == "COMPRAVENTA_INMUEBLE":
            data["municipio"] = ""
        elif code == "ARRENDAMIENTO":
            data["renta"] = ""
        elif code == "CONSTITUCION_SOCIEDAD":
            data["representante_legal"] = ""
        elif code == "DONACION_ENTRE_VIVOS":
            data["bienes"] = ""
        else:
            data["referencia_acta"] = ""
        return _expected(["RULE-001"], code)
    return 0


def seed_cases_and_documents(
    db: Session, deed_types: dict, templates: dict, users: dict, clients: list[Client]
) -> list[tuple[TestCase, Document]]:
    existing = db.query(TestCase).count()
    if existing:
        results: list[tuple[TestCase, Document]] = []
        for tc in db.query(TestCase).all():
            doc = db.get(Document, tc.document_id) if tc.document_id else None
            if doc is not None:
                results.append((tc, doc))
        return results

    responsable = users["notario.demo"]
    created: list[tuple[TestCase, Document]] = []
    counter = 1
    case_seq = 1

    for code, _name, _desc in DEED_TYPES:
        template = templates[code]
        version = next(v for v in template.versions if v.is_active)
        for n in range(20):
            kind = INCONSISTENCY_TYPES[n % len(INCONSISTENCY_TYPES)]
            data = build_field_data(code, db, clients, counter)
            expected = apply_inconsistency(code, data, kind)

            client = random.choice(clients)
            numero = f"2026-{case_seq:04d}"
            case_seq += 1
            case = Case(
                numero_expediente=numero,
                client_id=client.id,
                deed_type_id=deed_types[code].id,
                responsable_id=responsable.id,
                estado="ABIERTO",
                fecha_apertura=datetime(2026, 1, 1, tzinfo=UTC) + timedelta(days=counter),
                observaciones="Expediente sintético de prueba.",
            )
            db.add(case)
            db.flush()

            doc = Document(
                case_id=case.id,
                template_version_id=version.id,
                title=f"Borrador {code} #{counter:03d}",
                status="BORRADOR",
                field_data=data,
                created_by=responsable.id,
            )
            db.add(doc)
            db.flush()

            tc = TestCase(
                code=f"TC-{counter:03d}",
                deed_type_code=code,
                title=f"Caso sintético {counter:03d} ({kind})",
                description=f"Caso con condición: {kind}",
                inconsistency_types=kind,
                expected_errors=expected,
                document_id=doc.id,
            )
            db.add(tc)
            db.flush()
            created.append((tc, doc))
            counter += 1
    db.flush()
    return created


def seed_measurements(db: Session, cases: list[tuple], user) -> None:
    if db.query(TestExecution).count():
        return
    for tc, _doc in cases:
        # Método tradicional: 200-260 minutos. Método con sistema: 50-80 minutos.
        trad = round(random.uniform(200.0, 260.0), 2)
        sysm = round(random.uniform(50.0, 80.0), 2)

        for method, minutes in (("TRADITIONAL", trad), ("SYSTEM", sysm)):
            execution = TestExecution(
                test_case_id=tc.id,
                method=method,
                executed_by=user.id,
                notes="Medición sintética para el experimento de tesis.",
            )
            db.add(execution)
            db.flush()
            start = datetime(2026, 6, 1, 9, 0, tzinfo=UTC)
            db.add(
                TimeMeasurement(
                    execution_id=execution.id,
                    started_at=start,
                    finished_at=start + timedelta(minutes=minutes),
                    duration_seconds=minutes * 60.0,
                    duration_minutes=minutes,
                    errors_found=tc.expected_errors
                    if method == "SYSTEM"
                    else max(0, tc.expected_errors - 1),
                    errors_missed=0 if method == "SYSTEM" else 1,
                    corrections=tc.expected_errors,
                    notes="Medición sintética.",
                )
            )
    db.flush()


def seed_billing(db: Session, cases: list[tuple], users: dict) -> None:
    if db.query(Quote).count():
        return

    user: User = users["adminfin.demo"]
    tc, doc = cases[0]
    case = db.get(Case, doc.case_id)
    assert case is not None

    quote = Quote(
        numero="COT-2026-0001",
        client_id=case.client_id,
        case_id=case.id,
        fecha=date(2026, 3, 1),
        fecha_vencimiento=date(2026, 3, 31),
        estado="ENVIADA",
        descuento=Decimal("100.00"),
        observaciones="Cotización sintética.",
        created_by=user.id,
    )
    db.add(quote)
    db.flush()
    for i, (desc, precio) in enumerate(
        [("Honorarios de escritura", Decimal("2500.00")), ("Gastos notariales", Decimal("750.00"))]
    ):
        item = QuoteItem(
            quote_id=quote.id,
            descripcion=desc,
            cantidad=Decimal("1"),
            precio_unitario=precio,
            order_index=i,
        )
        item.recalculate()
        db.add(item)
    db.flush()
    db.refresh(quote)
    quote.recalculate()

    budget = Budget(
        case_id=case.id,
        nombre="Presupuesto sintético del expediente",
        created_by=user.id,
        observaciones="Presupuesto de ejemplo.",
    )
    db.add(budget)
    db.flush()
    for i, (tipo, desc, monto) in enumerate(
        [
            ("INGRESO", "Honorarios", Decimal("4000.00")),
            ("GASTO", "Papelería", Decimal("300.00")),
            ("OTRO", "Viáticos", Decimal("150.00")),
        ]
    ):
        db.add(
            BudgetItem(budget_id=budget.id, tipo=tipo, descripcion=desc, monto=monto, order_index=i)
        )
    db.flush()
    db.refresh(budget)
    budget.recalculate()

    charge = Charge(
        client_id=case.client_id,
        case_id=case.id,
        concepto="Cobro por servicios notariales",
        monto=Decimal("3250.00"),
        monto_pagado=Decimal("0.00"),
        saldo=Decimal("3250.00"),
        fecha_emision=date(2026, 4, 1),
        fecha_vencimiento=date(2026, 4, 30),
        estado="PENDIENTE",
        created_by=user.id,
    )
    db.add(charge)
    db.flush()
    payment = Payment(
        charge_id=charge.id,
        fecha=date(2026, 4, 15),
        monto=Decimal("1000.00"),
        medio_pago="TRANSFERENCIA",
        referencia="REF-SINTETICA-001",
        created_by=user.id,
    )
    db.add(payment)
    db.flush()
    db.refresh(charge)
    charge.recalculate()
    db.flush()


def run(reset: bool = False) -> None:
    ensure_storage_dirs()
    if reset:
        Base.metadata.drop_all(bind=engine)
    Base.metadata.create_all(bind=engine)

    db: Session = SessionLocal()
    try:
        from app.models.user import User

        refs = seed_reference_data(db)
        users = seed_users(db, refs["roles"])
        seed_validation_rules(db)
        templates = seed_templates(db, refs["deed_types"], users)
        clients = seed_clients(db)
        cases = seed_cases_and_documents(db, refs["deed_types"], templates, users, clients)
        seed_measurements(db, cases, users["notario.demo"])
        seed_billing(db, cases, users)
        db.commit()

        admin: User = users["admin"]
        print("Seed completado.")
        print(f"  Usuarios: {db.query(User).count()}")
        print(f"  Clientes: {db.query(Client).count()}")
        print(f"  Expedientes: {db.query(Case).count()}")
        print(f"  Plantillas: {db.query(Template).count()}")
        print(f"  Documentos: {db.query(Document).count()}")
        print(f"  Casos de prueba: {db.query(TestCase).count()}")
        print(f"  Reglas de validación: {db.query(ValidationRule).count()}")
        print(f"  Mediciones: {db.query(TimeMeasurement).count()}")
        print(f"  Administrador: {admin.username}")
    finally:
        db.close()


def main() -> None:
    parser = argparse.ArgumentParser(description="Genera datos sintéticos del sistema.")
    parser.add_argument("--reset", action="store_true", help="Elimina y recrea todas las tablas.")
    args = parser.parse_args()
    run(reset=args.reset)


if __name__ == "__main__":
    main()
