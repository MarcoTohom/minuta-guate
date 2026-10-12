# backend/tests/integration/test_documents.py
"""Pruebas de integración de borradores: validación y generación DOCX."""

from __future__ import annotations

from pathlib import Path

from docx import Document as DocxDocument

from tests.conftest import make_docx

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"

VALID_COMPRAVENTA = {
    "numero_escritura": 1,
    "fecha_escritura": "2026-03-15",
    "comprador_nombre": "Juan Pérez López",
    "comprador_dpi": "1234567890101",
    "comprador_nit": "12345678-9",
    "vendedor_nombre": "María López Pérez",
    "vendedor_dpi": "1234567890102",
    "precio": "100000.00",
    "finca": "12345",
    "folio": "100",
    "libro": "5",
    "departamento": "Guatemala",
    "municipio": "Mixco",
}

TEMPLATE_BODY = [
    "ESCRITURA PÚBLICA NÚMERO {{ numero_escritura }}.",
    "Comprador {{ comprador_nombre }} con DPI {{ comprador_dpi }}.",
    "Precio {{ precio }}.",
]


def _deed_type_id(client, headers, code="COMPRAVENTA_INMUEBLE") -> str:
    response = client.get("/api/v1/deed-types", headers=headers)
    return next(d for d in response.json() if d["code"] == code)["id"]


def setup_comparaventa(client, headers) -> dict:
    """Crea cliente, expediente, plantilla con versión/campos y devuelve ids."""
    client_id = client.post(
        "/api/v1/clients",
        json={
            "tipo_cliente": "INDIVIDUAL",
            "person": {"nombres": "Juan", "apellidos": "Pérez López", "dpi": "1234567890101"},
        },
        headers=headers,
    ).json()["id"]

    case_id = client.post(
        "/api/v1/cases",
        json={"client_id": client_id, "deed_type_id": _deed_type_id(client, headers)},
        headers=headers,
    ).json()["id"]

    template_id = client.post(
        "/api/v1/templates",
        json={"name": "Compraventa integración", "deed_type_id": _deed_type_id(client, headers)},
        headers=headers,
    ).json()["id"]

    files = {"file": ("compraventa.docx", make_docx(TEMPLATE_BODY), DOCX_MIME)}
    version = client.post(
        f"/api/v1/templates/{template_id}/versions", files=files, headers=headers
    ).json()
    client.post(
        f"/api/v1/templates/{template_id}/activate",
        json={"version_id": version["id"]},
        headers=headers,
    )
    client.post(
        f"/api/v1/templates/{template_id}/versions/{version['id']}/fields",
        json={"variables": ["numero_escritura", "comprador_nombre", "comprador_dpi", "precio"]},
        headers=headers,
    )
    return {
        "client_id": client_id,
        "case_id": case_id,
        "template_id": template_id,
        "template_version_id": version["id"],
    }


def create_document(client, headers, case_id, version_id, field_data, title="Borrador"):
    return client.post(
        "/api/v1/documents",
        json={
            "case_id": case_id,
            "template_version_id": version_id,
            "title": title,
            "field_data": field_data,
        },
        headers=headers,
    )


class TestDocumentCrud:
    def test_create_and_get(self, client, auth):
        ctx = setup_comparaventa(client, auth["notario"])
        response = create_document(
            client, auth["notario"], ctx["case_id"], ctx["template_version_id"], VALID_COMPRAVENTA
        )
        assert response.status_code == 201, response.text
        document_id = response.json()["id"]
        got = client.get(f"/api/v1/documents/{document_id}", headers=auth["notario"])
        assert got.status_code == 200
        assert got.json()["title"] == "Borrador"

    def test_update_and_list(self, client, auth):
        ctx = setup_comparaventa(client, auth["notario"])
        document_id = create_document(
            client, auth["notario"], ctx["case_id"], ctx["template_version_id"], VALID_COMPRAVENTA
        ).json()["id"]
        updated = client.put(
            f"/api/v1/documents/{document_id}",
            json={"title": "Borrador actualizado"},
            headers=auth["notario"],
        )
        assert updated.status_code == 200
        listing = client.get(
            "/api/v1/documents", params={"case_id": ctx["case_id"]}, headers=auth["notario"]
        )
        assert listing.json()["total"] >= 1

    def test_create_requires_permission(self, client, auth):
        ctx = setup_comparaventa(client, auth["notario"])
        response = create_document(
            client, auth["adminfin"], ctx["case_id"], ctx["template_version_id"], {}
        )
        assert response.status_code == 403


class TestValidation:
    def test_valid_document_can_generate(self, client, auth):
        ctx = setup_comparaventa(client, auth["notario"])
        document_id = create_document(
            client, auth["notario"], ctx["case_id"], ctx["template_version_id"], VALID_COMPRAVENTA
        ).json()["id"]
        response = client.post(f"/api/v1/documents/{document_id}/validate", headers=auth["notario"])
        assert response.status_code == 200, response.text
        summary = response.json()
        assert summary["critical"] == 0
        assert summary["error"] == 0
        assert summary["can_generate"] is True

    def test_inconsistent_dpi_is_critical(self, client, auth):
        ctx = setup_comparaventa(client, auth["notario"])
        data = dict(VALID_COMPRAVENTA, expediente_comprador_dpi="9999999999999")
        document_id = create_document(
            client, auth["notario"], ctx["case_id"], ctx["template_version_id"], data
        ).json()["id"]
        summary = client.post(
            f"/api/v1/documents/{document_id}/validate", headers=auth["notario"]
        ).json()
        assert summary["critical"] >= 1
        assert summary["can_generate"] is False
        assert any(r["rule_id"] == "RULE-004" for r in summary["results"])

    def test_missing_required_field_is_error(self, client, auth):
        ctx = setup_comparaventa(client, auth["notario"])
        data = dict(VALID_COMPRAVENTA)
        data["vendedor_nombre"] = ""
        document_id = create_document(
            client, auth["notario"], ctx["case_id"], ctx["template_version_id"], data
        ).json()["id"]
        summary = client.post(
            f"/api/v1/documents/{document_id}/validate", headers=auth["notario"]
        ).json()
        assert any(r["rule_id"] == "RULE-001" for r in summary["results"])
        assert summary["can_generate"] is False

    def test_review_validation_result(self, client, auth):
        ctx = setup_comparaventa(client, auth["notario"])
        data = dict(VALID_COMPRAVENTA, expediente_comprador_dpi="9999999999999")
        document_id = create_document(
            client, auth["notario"], ctx["case_id"], ctx["template_version_id"], data
        ).json()["id"]
        summary = client.post(
            f"/api/v1/documents/{document_id}/validate", headers=auth["notario"]
        ).json()
        result_id = summary["results"][0]["id"]
        response = client.post(
            f"/api/v1/documents/{document_id}/validations/{result_id}/review",
            json={"reviewed": True, "correction_note": "Corregido manualmente"},
            headers=auth["notario"],
        )
        assert response.status_code == 200
        assert response.json()["reviewed"] is True

    def test_validations_endpoint(self, client, auth):
        ctx = setup_comparaventa(client, auth["notario"])
        document_id = create_document(
            client, auth["notario"], ctx["case_id"], ctx["template_version_id"], VALID_COMPRAVENTA
        ).json()["id"]
        client.post(f"/api/v1/documents/{document_id}/validate", headers=auth["notario"])
        response = client.get(f"/api/v1/validations/{document_id}", headers=auth["notario"])
        assert response.status_code == 200
        assert "results" in response.json()


class TestGeneration:
    def test_generate_docx(self, client, auth):
        ctx = setup_comparaventa(client, auth["notario"])
        document_id = create_document(
            client, auth["notario"], ctx["case_id"], ctx["template_version_id"], VALID_COMPRAVENTA
        ).json()["id"]
        client.post(f"/api/v1/documents/{document_id}/validate", headers=auth["notario"])

        response = client.post(f"/api/v1/documents/{document_id}/generate", headers=auth["notario"])
        assert response.status_code == 200, response.text
        version = response.json()
        assert version["version_number"] == 1
        assert version["file_path"]

        # El archivo existe, es DOCX, se puede abrir y no quedan placeholders.
        path = Path(version["file_path"])
        assert path.exists()
        assert path.suffix == ".docx"
        opened = DocxDocument(str(path))
        text = "\n".join(p.text for p in opened.paragraphs)
        assert "{{" not in text
        assert "1234567890101" in text  # dato sustituido presente

    def test_generation_blocked_on_errors(self, client, auth):
        ctx = setup_comparaventa(client, auth["notario"])
        data = dict(VALID_COMPRAVENTA, expediente_comprador_dpi="9999999999999")
        document_id = create_document(
            client, auth["notario"], ctx["case_id"], ctx["template_version_id"], data
        ).json()["id"]
        client.post(f"/api/v1/documents/{document_id}/validate", headers=auth["notario"])
        response = client.post(f"/api/v1/documents/{document_id}/generate", headers=auth["notario"])
        assert response.status_code == 422

    def test_versions_are_not_overwritten(self, client, auth):
        ctx = setup_comparaventa(client, auth["notario"])
        document_id = create_document(
            client, auth["notario"], ctx["case_id"], ctx["template_version_id"], VALID_COMPRAVENTA
        ).json()["id"]
        client.post(f"/api/v1/documents/{document_id}/validate", headers=auth["notario"])
        first = client.post(
            f"/api/v1/documents/{document_id}/generate", headers=auth["notario"]
        ).json()
        second = client.post(
            f"/api/v1/documents/{document_id}/generate", headers=auth["notario"]
        ).json()
        assert first["version_number"] == 1
        assert second["version_number"] == 2
        assert first["file_path"] != second["file_path"]

        versions = client.get(
            f"/api/v1/documents/{document_id}/versions", headers=auth["notario"]
        ).json()
        assert len(versions) == 2

    def test_download_version(self, client, auth):
        ctx = setup_comparaventa(client, auth["notario"])
        document_id = create_document(
            client, auth["notario"], ctx["case_id"], ctx["template_version_id"], VALID_COMPRAVENTA
        ).json()["id"]
        client.post(f"/api/v1/documents/{document_id}/validate", headers=auth["notario"])
        version = client.post(
            f"/api/v1/documents/{document_id}/generate", headers=auth["notario"]
        ).json()
        response = client.get(
            f"/api/v1/documents/{document_id}/versions/{version['id']}/download",
            headers=auth["notario"],
        )
        assert response.status_code == 200
        assert response.content.startswith(b"PK")

    def test_generate_requires_permission(self, client, auth):
        ctx = setup_comparaventa(client, auth["notario"])
        document_id = create_document(
            client, auth["notario"], ctx["case_id"], ctx["template_version_id"], VALID_COMPRAVENTA
        ).json()["id"]
        # AUXILIAR puede validar pero no generar.
        response = client.post(
            f"/api/v1/documents/{document_id}/generate", headers=auth["auxiliar"]
        )
        assert response.status_code == 403
